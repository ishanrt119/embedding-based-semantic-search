import tempfile
import os
from typing import List, Dict, Any,Optional
from retrieval_engine.retrieval_service import retrieval_service
from rag.rag_service import rag_service
from database.repositories.document_repository import DocumentRepository

class ReportGenerator:
    @staticmethod
    async def generate_markdown_report(report_type: str, doc_ids: List[str], user_id: str, topic: Optional[str] = None) -> str:
        all_contexts = []
        doc_names = []
        
        search_query = topic if topic and topic.strip() else "summary overview key findings important information"
        
        for doc_id in doc_ids:
            doc = await DocumentRepository.get_document_by_id(doc_id)
            if doc:
                doc_names.append(doc.get("name", "Unknown Document"))
                
            res = await retrieval_service.get_context(
                query=search_query,
                user_id=user_id,
                top_k=8,
                max_tokens=2000,
                dataset_id=doc_id
            )
            all_contexts.extend(res.get("contexts", []))
            
        all_contexts = all_contexts[:30]
        
        context_str = ""
        for idx, ctx in enumerate(all_contexts):
            doc_name = ctx.get("document_name", "Unknown")
            page = ctx.get("page_number", "?")
            context_str += f"\n--- Source {idx+1} ---\nDocument: {doc_name} (Page {page})\n{ctx.get('content')}\n"
            
        if topic and topic.strip():
            prompt = f"""You are an expert Report Writer.
Generate a comprehensive '{report_type}' regarding the topic: '{topic}'.
Use the following documents as your exclusive source material: {', '.join(doc_names)}.

CONTEXT:
{context_str}

RULES:
1. Structure the report beautifully using Markdown.
2. Include an Executive Summary, Key Themes, Important Findings, and Action Items.
3. You MUST cite your sources using [Source X] inline.
4. Do NOT hallucinate information.
"""
        else:
            prompt = f"""You are an expert Report Writer.
Generate a comprehensive '{report_type}' automatically based on the provided document excerpts.
Use the document titles and content to infer the scope and key themes of the report.
Source documents: {', '.join(doc_names)}.

CONTEXT:
{context_str}

RULES:
1. Structure the report beautifully using Markdown.
2. You MUST include exactly these sections:
   - Executive Summary
   - Key Findings
   - Important Information
   - Action Items (if applicable)
   - Sources
3. You MUST cite your sources using [Source X] inline.
4. Do NOT hallucinate information. Use only the provided context.
"""

        try:
            raw_ans = await rag_service.generate_answer([{"role": "user", "content": prompt}])
            return raw_ans
        except Exception as e:
            return f"# Error Generating Report\n{str(e)}"
            
    @staticmethod
    def export_to_pdf(markdown_content: str) -> bytes:
        try:
            from fpdf import FPDF
            import markdown
            import tempfile
            import os
            
            pdf = FPDF()
            pdf.add_page()
            
            # Simple conversion of markdown to HTML for FPDF
            html = markdown.markdown(markdown_content)
            
            # Write HTML to PDF
            pdf.write_html(html)
            
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_pdf:
                pdf_path = tmp_pdf.name
                
            pdf.output(pdf_path)
            
            with open(pdf_path, "rb") as f:
                pdf_bytes = f.read()
                
            os.remove(pdf_path)
            return pdf_bytes
        except Exception as e:
            raise RuntimeError(f"PDF generation failed: {str(e)}")
            
    @staticmethod
    def export_to_docx(markdown_content: str) -> bytes:
        try:
            from docx import Document
            import tempfile
            import os
            
            doc = Document()
            doc.add_heading('Generated Report', 0)
            
            # Very basic markdown to docx conversion for MVP
            for line in markdown_content.split('\n'):
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('- '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.strip():
                    doc.add_paragraph(line)
                    
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_docx:
                docx_path = tmp_docx.name
                
            doc.save(docx_path)
            
            with open(docx_path, "rb") as f:
                docx_bytes = f.read()
                
            os.remove(docx_path)
            return docx_bytes
        except Exception as e:
            raise RuntimeError(f"DOCX generation failed: {str(e)}")

report_generator = ReportGenerator()
